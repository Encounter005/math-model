"""Build a validated source manifest from the supplied label workbook."""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

import yaml
from openpyxl import load_workbook

FIELDNAMES = (
    "id",
    "video_id",
    "clip_id",
    "text",
    "label",
    "annotation",
    "video_path",
)
REQUIRED_COLUMNS = ("video_id", "clip_id", "text", "label", "annotation")


def build_manifest(labels_path: Path, samples_root: Path, output_path: Path) -> None:
    """Validate workbook rows and write their resolved media locations as CSV."""
    workbook = load_workbook(labels_path, read_only=True, data_only=True)
    sheet = workbook.active
    columns = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    header = {str(value): index for index, value in enumerate(columns) if value is not None}
    missing_columns = set(REQUIRED_COLUMNS) - header.keys()
    if missing_columns:
        raise ValueError(f"missing label columns: {', '.join(sorted(missing_columns))}")

    rows: list[dict[str, str]] = []
    ids: set[str] = set()
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        video_id = str(values[header["video_id"]]).strip()
        clip_id = str(values[header["clip_id"]]).strip()
        text = values[header["text"]]
        sample_id = f"{video_id}_{clip_id}"
        if sample_id in ids:
            raise ValueError(f"row {row_number}: duplicate ID {sample_id}")
        ids.add(sample_id)
        if not isinstance(text, str) or not text.strip():
            raise ValueError(f"row {row_number}: empty transcription for {sample_id}")
        try:
            label = float(values[header["label"]])
        except (TypeError, ValueError) as error:
            raise ValueError(f"row {row_number}: invalid label for {sample_id}") from error
        if not -3 <= label <= 3:
            raise ValueError(f"row {row_number}: label outside [-3, 3] for {sample_id}")
        video_path = (samples_root / video_id / f"{clip_id}.mp4").resolve()
        if not video_path.is_file():
            raise ValueError(f"row {row_number}: missing video for {sample_id}: {video_path}")
        rows.append(
            {
                "id": sample_id,
                "video_id": video_id,
                "clip_id": clip_id,
                "text": text,
                "label": str(label),
                "annotation": str(values[header["annotation"]]),
                "video_path": str(video_path),
            }
        )
    workbook.close()

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(file, fieldnames=FIELDNAMES)
        writer.writeheader()
        writer.writerows(rows)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", type=Path, default=Path("configs/question_1.yaml"))
    parser.add_argument("--labels", type=Path)
    parser.add_argument("--samples-root", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    config = yaml.safe_load(args.config.read_text(encoding="utf-8"))
    paths = config["paths"]
    build_manifest(
        args.labels or Path(paths["labels"]),
        args.samples_root or Path(paths["raw_samples"]),
        args.output or Path(paths["output"]) / "manifest.csv",
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
